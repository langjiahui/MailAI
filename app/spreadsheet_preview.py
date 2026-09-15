"""Read-only, bounded Excel cell previews; never evaluate formulas or macros."""
import io
import zipfile

ROWS, COLS, SHEETS = 200, 40, 10


def preview_workbook(payload, suffix):
    sheets = []
    def text(value):
        return '' if value is None else str(value)[:1000]
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Expanded workbook too large')
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book.worksheets[:SHEETS]:
                rows = [[text(v) for v in row] for row in sheet.iter_rows(
                    max_row=min(sheet.max_row or ROWS, ROWS),
                    max_col=min(sheet.max_column or COLS, COLS), values_only=True)]
                sheets.append({'name': sheet.title, 'rows': rows})
        finally:
            book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(file_contents=payload, on_demand=True)
        try:
            for index in range(min(book.nsheets, SHEETS)):
                sheet = book.sheet_by_index(index)
                rows = []
                for r in range(min(sheet.nrows, ROWS)):
                    row = []
                    for c in range(min(sheet.ncols, COLS)):
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        row.append(text(value))
                    rows.append(row)
                sheets.append({'name': sheet.name, 'rows': rows})
        finally:
            book.release_resources()
    for sheet in sheets:
        while sheet['rows'] and not any(sheet['rows'][-1]):
            sheet['rows'].pop()
    return {'kind': 'spreadsheet', 'sheets': sheets,
            'message': '最多预览前 10 个工作表，每表前 200 行、40 列；公式显示已保存的结果，未缓存时留空。图表和完整格式请下载查看。'}

"""Markdown to Word (.docx) converter for README."""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import parse_xml


def set_cell_shading(cell, fill):
    """Set cell background color (fill as hex, e.g. 'F2F2F2')."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def parse_markdown(md: str) -> Document:
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block
        if stripped.startswith('```'):
            lang = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing ```
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # Table
        if '|' in stripped and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set('|-:| '):
            rows = []
            while i < len(lines) and '|' in lines[i].strip():
                rows.append(lines[i].strip())
                i += 1
            if len(rows) >= 2:
                cells = [cell.strip() for cell in rows[0].split('|')[1:-1]]
                table = doc.add_table(rows=1, cols=len(cells))
                table.style = 'Light Grid Accent 1'
                hdr_cells = table.rows[0].cells
                for j, text in enumerate(cells):
                    hdr_cells[j].text = text
                    for para in hdr_cells[j].paragraphs:
                        para.paragraph_format.space_before = Pt(2)
                        para.paragraph_format.space_after = Pt(2)
                for row in rows[2:]:
                    cells = [cell.strip() for cell in row.split('|')[1:-1]]
                    if not cells:
                        continue
                    row_cells = table.add_row().cells
                    for j, text in enumerate(cells):
                        row_cells[j].text = text if j < len(cells) else ''
                        for para in row_cells[j].paragraphs:
                            para.paragraph_format.space_before = Pt(2)
                            para.paragraph_format.space_after = Pt(2)
            doc.add_paragraph()
            continue

        # Headings
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            p = doc.add_heading(text, level=min(level, 3))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        # Horizontal rule
        if stripped == '---' or stripped == '***':
            doc.add_paragraph().add_run().add_break()
            i += 1
            continue

        # Blockquote
        if stripped.startswith('>'):
            text = stripped[1:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(text)
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # List item
        if re.match(r'^[-*+]\s', stripped):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(text)
            i += 1
            continue

        if re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.add_run(text)
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Normal paragraph with simple inline formatting
        p = doc.add_paragraph()
        add_inline_runs(p, line)
        i += 1

    return doc


def add_inline_runs(paragraph, text):
    """Parse inline bold, italic, code and add runs."""
    # Simple regex-based parser for **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        chunk = m.group(0)
        if chunk.startswith('**') and chunk.endswith('**'):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith('*') and chunk.endswith('*'):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        elif chunk.startswith('`') and chunk.endswith('`'):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def main():
    parser = argparse.ArgumentParser(description='Convert Markdown README to Word document')
    parser.add_argument('input', help='Input markdown file')
    parser.add_argument('-o', '--output', help='Output docx file', default='README.docx')
    args = parser.parse_args()

    md = Path(args.input).read_text(encoding='utf-8')
    doc = parse_markdown(md)
    doc.save(args.output)
    print(f'Saved Word document to {args.output}')


if __name__ == '__main__':
    main()

"""Fill a real attached XLS/XLSX form and hand back a reply draft, never send."""
import io
import json
import os
import sys
import tempfile
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import assistant_actions, assistant_attachments, assistant_document_reply as task, config, db, mail_assistant


def _xlsx():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = '登记表'
    sheet['A1'] = '姓名'
    sheet['A2'] = '车牌号'
    sheet['B1'].number_format = '@'
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _xls():
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('登记表')
    sheet.write(0, 0, '姓名')
    sheet.write(1, 0, '车牌号')
    sheet.write(0, 1, '', xlwt.easyxf('borders: bottom thin;'))
    sheet.set_portrait(0)
    sheet.set_left_margin(0.16)
    sheet.set_right_margin(0.2)
    sheet.set_header_str(b'')
    sheet.set_footer_str(b'')
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _formula_xls():
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('登记表')
    sheet.write(0, 0, '姓名')
    sheet.write(0, 2, xlwt.Formula('1+2'))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _tabular_xls():
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet('登记表')
    for col, value in enumerate(['部门','工号','姓名','手机']):
        sheet.write(0, col, value)
    sheet.write(1, 0, 'ERP部')
    sheet.write(2, 0, 'EAM部')
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _email(root, name, payload, uid):
    message = EmailMessage()
    message['From'] = 'sender@example.com'
    message['To'] = 'me@example.com'
    message['Subject'] = '假期自驾登记'
    message['Message-ID'] = '<source@example.com>'
    message.set_content('请填写附件登记表后回复。')
    message.add_attachment(payload, maintype='application', subtype='octet-stream', filename=name)
    path = Path(root) / f'source-{name}.eml'
    path.write_bytes(message.as_bytes())
    with db.conn() as c:
        cur = c.execute("INSERT INTO emails(uid,folder,from_addr,to_addr,subject,body_text,date,message_id,raw_path,attachments,created_at,remote_missing) "
                        "VALUES(?,?,?,?,?,?,?,?,?,?,?,0)",
                        (uid, 'INBOX', 'sender@example.com', 'me@example.com', '假期自驾登记',
                         '请填写附件登记表后回复。', '2026-09-29T10:00:00', '<source@example.com>',
                         str(path), json.dumps([{'name':name,'size':len(payload)}]), '2026-09-29T10:00:00'))
        return cur.lastrowid


def _model(messages, **kwargs):
    if '只输出 JSON 对象' in messages[0]['content']:
        if '仅从用户本次' in messages[0]['content']:
            value = {'values':['张三','沪A12345']}
        else:
            value = {'fields':[{'sheet':'登记表','cell':'B1','label':'姓名'},
                               {'sheet':'登记表','cell':'B2','label':'车牌号'}]}
    else:
        value = '您好，登记表已填写，请查收附件。'
    return {'choices':[{'message':{'content':json.dumps(value, ensure_ascii=False) if isinstance(value,dict) else value}}]}


def _run_one(root, ext, payload):
    email_id = _email(root, f'登记表{ext}', payload, 1 if ext == '.xlsx' else 2)
    proposal = assistant_actions.detect_proposal('帮我填写附件并回复这封邮件', [email_id])
    assert proposal and proposal['type'] == 'fill_attachment_reply'
    assert mail_assistant.ask('帮我填写附件并回复这封邮件', email_ids=[email_id])['action']['type'] == 'fill_attachment_reply'
    assert ('action', proposal) in list(mail_assistant.ask_stream('帮我填写附件并回复这封邮件', email_ids=[email_id]))
    # The short wording in the inbox must use the document workflow even when
    # cached attachment metadata has not caught up with the original message.
    original_attachments = db.get_email(email_id)['attachments']
    original_draft_count = len(db.list_drafts())
    db.update_attachment_metadata(email_id, [])
    short = '填写附件回复邮件'
    for wording in (short, '填好文档回邮件', '回复邮件时填写附件'):
        assert assistant_actions.requests_document_reply(wording), wording
    catalog = assistant_attachments.catalog(email_id)
    assert catalog['items'][0]['name'] == f'登记表{ext}' and catalog['items'][0]['supported']
    assert assistant_actions.detect_proposal(short, [email_id])['type'] == 'fill_attachment_reply'
    assert mail_assistant.ask(short, email_ids=[email_id])['action']['type'] == 'fill_attachment_reply'
    assert any(event == 'action' and value['type'] == 'fill_attachment_reply'
               for event, value in mail_assistant.ask_stream(short, email_ids=[email_id]))
    assert not assistant_actions.detect_proposal(short)
    no_scope = mail_assistant.ask(short)
    assert '请先打开' in no_scope['answer'] and 'action' not in no_scope
    no_scope_stream = list(mail_assistant.ask_stream(short))
    assert not any(event == 'action' for event, _ in no_scope_stream)
    assert any(event == 'delta' and '请先打开' in value for event, value in no_scope_stream)
    many = mail_assistant.ask(short, email_ids=[email_id, email_id + 1])
    assert '请只选择一封' in many['answer'] and 'action' not in many
    assert len(db.list_drafts()) == original_draft_count, 'a fill request without one source must never create a plain reply draft'
    db.update_attachment_metadata(email_id, original_attachments)
    with patch.object(task.client, 'chat_completion', side_effect=_model):
        plan = task.plan(email_id, 0)
        try:
            task.prepare(email_id, 0, '0'*64, [dict(item, value='测试') for item in plan['fields']], plan_token=plan['plan_token'])
            raise AssertionError('过期附件摘要应失败')
        except ValueError:
            pass
        assert [item['cell'] for item in plan['fields']] == ['B1','B2']
        extracted = task.suggest_values(email_id, 0, plan['digest'], plan['fields'], plan['plan_token'], '张三，车牌沪A12345')
        assert extracted['values'] == ['张三','沪A12345']
        fields = [dict(item, value=value) for item, value in zip(plan['fields'], ['张三','沪A12345'])]
        try:
            task.prepare(email_id, 0, plan['digest'], [dict(fields[0], cell='C1'),fields[1]], plan_token=plan['plan_token'])
            raise AssertionError('篡改计划中的目标单元格应失败')
        except ValueError as exc:
            assert '修改' in str(exc)
        result = task.prepare(email_id, 0, plan['digest'], fields, plan_token=plan['plan_token'])
    draft = db.get_draft(result['draft_id'])
    assert draft['to_addr'] == 'sender@example.com'
    assert draft['mode'] == 'reply' and draft['reply_to_email_id'] == email_id
    assert len(draft['attachments']) == 1
    generated = db.get_sent_attachment(result['draft_id'], 0, draft=True)
    assert generated['name'].endswith(f'_已填写{ext}')
    if ext == '.xls':
        original_print = task._legacy_print_settings(payload, 1)[0]
        generated_print = task._legacy_print_settings(generated['payload'], 1)[0]
        assert original_print['portrait'] == generated_print['portrait'] == 0
        assert original_print['margins'] == generated_print['margins']
        assert original_print['blank_header'] == generated_print['blank_header']
    book = task._book(generated['payload'], ext)
    try:
        sheet = dict((title, sheet) for title,sheet,_,_ in task._sheets(book, ext))['登记表']
        assert task._value(sheet, 'B1', ext) == '张三'
        assert task._value(sheet, 'B2', ext) == '沪A12345'
        assert task._value(sheet, 'A1', ext) == '姓名'
        if ext == '.xls':
            assert book.xf_list[sheet.cell(0, 1).xf_index].border.bottom_line_style == 1
    finally:
        if ext == '.xls': book.release_resources()
    assert not db.list_sent_messages(), 'prepare must never send'
    assert task._source(email_id, 0)[3] == payload, 'source attachment must remain unchanged'


def test_document_reply():
    with tempfile.TemporaryDirectory() as root, \
         patch.object(config, 'DB_PATH', os.path.join(root, 'mail.db')), \
         patch.object(config, 'IMAP_USER', 'me@example.com'):
        db.init_db()
        _run_one(root, '.xlsx', _xlsx())
        _run_one(root, '.xls', _xls())
        table_id = _email(root, '部门登记表.xls', _tabular_xls(), 3)
        choices = task.plan(table_id, 0)
        assert choices['needs_row_choice'] and len(choices['choices']) == 2
        with patch.object(task.client, 'chat_completion', side_effect=_model):
            selected = task.plan(table_id, 0, row_choice=3)
            assert [item['cell'] for item in selected['fields']] == ['B3','C3','D3']
            filled_fields = [dict(item, value=value) for item, value in zip(selected['fields'], ['007','张三','13800000000'])]
            result = task.prepare(table_id, 0, selected['digest'], filled_fields, plan_token=selected['plan_token'])
        generated = db.get_sent_attachment(result['draft_id'], 0, draft=True)
        book = task._book(generated['payload'], '.xls')
        try:
            sheet = book.sheet_by_index(0)
            assert sheet.cell_value(2, 1) == '007' and sheet.cell_value(2, 2) == '张三'
            assert sheet.cell_value(1, 1) == '', 'another department row must stay empty'
        finally: book.release_resources()
        literal = task._filled(_xlsx(), '.xlsx', [{'sheet':'登记表','cell':'B1','value':'=1+1'}])
        book = task._book(literal, '.xlsx')
        assert book['登记表']['B1'].value == '=1+1' and book['登记表']['B1'].data_type == 's'
        try:
            task._check_legacy_features(_formula_xls())
            raise AssertionError('带公式的旧版 XLS 不应重写')
        except ValueError as exc:
            assert '公式' in str(exc)


if __name__ == '__main__':
    test_document_reply()
    print('✅ 小邮填写 XLS/XLSX 并生成待审核回复草稿测试通过')

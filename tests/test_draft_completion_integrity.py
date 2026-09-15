"""No network: retain failed drafts, clean successful source drafts, preserve binary MIME."""
import base64, hashlib, sys, tempfile, json
from pathlib import Path
from unittest.mock import patch
from email import policy
from email.parser import BytesParser
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import db, config, smtp_client, outbox
from app.web import server

def main():
 with tempfile.TemporaryDirectory() as root, patch.object(config,'DB_PATH',str(Path(root)/'mail.db')), patch.object(config,'IMAP_USER','me@example.test'), patch.object(smtp_client,'configured',return_value=True), patch.object(smtp_client,'identity_matches_current_mailbox',return_value=True):
  db.init_db()
  source=db.upsert_email({'uid':9,'folder':'Drafts','status':'draft','subject':'source'})
  raw=bytes(range(256))*257+b'\x00\xff\r\n'
  attachment={'filename':'中文附件.xlsx','size':len(raw),'data_base64':base64.b64encode(raw).decode(),'sha256':hashlib.sha256(raw).hexdigest(),'content_type':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}
  data={'to_addr':'you@example.test','subject':'test','body_html':'<p>body</p>','attachments':[attachment],'source_draft_email_id':source,'preflight_confirmed':True}
  draft=db.save_draft(data);db.save_draft(data,draft)
  saved=db.get_draft(draft)
  message,_=smtp_client.build_message(saved)
  parsed=BytesParser(policy=policy.default).parsebytes(message.as_bytes(policy=policy.SMTP))
  part=next(parsed.iter_attachments());assert part.get_payload(decode=True)==raw;assert part.get_filename()==attachment['filename']
  from openpyxl import Workbook, load_workbook
  from io import BytesIO
  workbook=Workbook();workbook.active['A1']='中文附件测试';buffer=BytesIO();workbook.save(buffer)
  excel=buffer.getvalue()
  excel_att={**attachment,'data_base64':base64.b64encode(excel).decode(),'size':len(excel),'sha256':hashlib.sha256(excel).hexdigest()}
  excel_draft=db.save_draft({**data,'source_draft_email_id':None,'attachments':[excel_att]})
  excel_msg,_=smtp_client.build_message(db.get_draft(excel_draft))
  excel_received=next(BytesParser(policy=policy.default).parsebytes(excel_msg.as_bytes(policy=policy.SMTP)).iter_attachments()).get_payload(decode=True)
  assert excel_received==excel
  assert load_workbook(BytesIO(excel_received)).active['A1'].value=='中文附件测试'
  db.delete_draft(excel_draft)

  for bad in [dict(attachment,data_base64=''),dict(attachment,data_base64=base64.b64encode(raw[:-3]).decode()),dict(attachment,sha256='0'*64),{'filename':'missing.txt'}]:
   try:smtp_client.build_message({**data,'attachments':[bad]})
   except ValueError:pass
   else:raise AssertionError('Damaged attachment must be blocked')
  with patch.object(server.smtp_client,'send',side_effect=ValueError('refused')),patch.object(server.outgoing_guard,'local_issues',return_value=[]):
   try:server.api_send_mail(server.SendMailRequest(**data,id=draft))
   except Exception:pass
  assert db.get_draft(draft) and not db.get_email(source)['pending_action']
  with patch.object(server.smtp_client,'send',return_value={'sent_folder':'Sent','message_id':'<sent>','recipients':['you@example.test']}),patch.object(server.outgoing_guard,'local_issues',return_value=[]):
   assert server.api_send_mail(server.SendMailRequest(**data,id=draft))['ok']
  assert db.get_draft(draft) is None
  try:server.api_save_draft(server.DraftRequest(**data,id=draft))
  except Exception as error:assert getattr(error,'status_code',None)==409
  else:raise AssertionError('A late autosave must not recreate a sent draft')
  assert db.get_email(source)['pending_action']=='trash'
  assert db.get_email(source)['remote_missing']
  assert not server.api_drafts()
  second=db.save_draft(data)
  outbox.enqueue('unknown',{'id':second})
  with db.conn() as c:c.execute("UPDATE outbox SET status='unknown' WHERE token='unknown'")
  outbox.resolve('unknown',False);assert db.get_draft(second)
 print('PASS binary MIME integrity, corrupt payload rejection, successful draft cleanup and failure retention')
if __name__=='__main__':main()
